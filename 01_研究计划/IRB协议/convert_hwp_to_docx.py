import os
import re
import sys
import time
import zlib
from pathlib import Path

sys.path.insert(0, "/private/tmp/pydeps")

import olefile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path("5. 심의용 연구계획서_신진과제_ver2.1.hwp")
ORIGINAL_DOCX = Path("5. 심의용 연구계획서_신진과제_ver2.1_원문.docx")
CHINESE_DOCX = Path("5. 심의용 연구계획서_신진과제_ver2.1_中文.docx")
EXTRACTED_TXT = Path("5. 심의용 연구계획서_신진과제_ver2.1_원문추출.txt")
TRANSLATION_CACHE = Path("5. 심의용 연구계획서_신진과제_ver2.1_中文段落.txt")


def set_font(run, east_asia="Malgun Gothic", latin="Arial", size=10.5, bold=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def clean_text(text):
    text = text.replace("\u000b", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\uffff]", "", text)
    # HWP inline control payloads can decode as unrelated ideographs or exotic glyphs.
    text = re.sub(r"[\u3400-\u4dbf\u4e00-\u9fff\u0f00-\u0fff\u0b80-\u0bff]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_hwp_paragraphs(path):
    ole = olefile.OleFileIO(str(path))
    flags = int.from_bytes(ole.openstream("FileHeader").read()[36:40], "little")
    compressed = bool(flags & 1)
    paragraphs = []
    section_names = sorted(
        "/".join(item) for item in ole.listdir() if len(item) == 2 and item[0] == "BodyText"
    )
    for name in section_names:
        raw = ole.openstream(name).read()
        data = zlib.decompress(raw, -15) if compressed else raw
        i = 0
        while i + 4 <= len(data):
            header = int.from_bytes(data[i : i + 4], "little")
            i += 4
            tag_id = header & 0x3FF
            size = (header >> 20) & 0xFFF
            if size == 0xFFF:
                size = int.from_bytes(data[i : i + 4], "little")
                i += 4
            payload = data[i : i + size]
            i += size
            if tag_id == 67:
                text = clean_text(payload.decode("utf-16le", "ignore"))
                if text:
                    paragraphs.append(text)
    return paragraphs


def paragraph_kind(text, index):
    if index == 0 or text in {"(인간대상 연구용)", "(ver. 2.1)"}:
        return "title"
    if re.match(r"^\d+(\.\d+)*\.\s+", text) or re.match(r"^\d+\.\s*[\w가-힣]", text):
        return "heading"
    if re.match(r"^\(?[가-힣a-z]\)|^[a-z]\.\s|^[ⅰⅱⅲⅳⅴ]\.", text):
        return "list"
    return "normal"


def add_paragraph(doc, text, kind="normal", east_asia="Malgun Gothic"):
    style = "Normal"
    if kind == "heading":
        style = "Heading 1" if re.match(r"^\d+\.\s", text) else "Heading 2"
    p = doc.add_paragraph(style=style)
    if kind == "title":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_font(
        run,
        east_asia=east_asia,
        latin="Arial",
        size=15 if kind == "title" else 11 if kind == "heading" else 10,
        bold=kind in {"title", "heading"},
    )


def write_docx(paragraphs, path, language="ko"):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(56.7)
    section.bottom_margin = Pt(56.7)
    section.left_margin = Pt(56.7)
    section.right_margin = Pt(56.7)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei" if language == "zh" else "Malgun Gothic")
    for i, para in enumerate(paragraphs):
        add_paragraph(doc, para, paragraph_kind(para, i), "Microsoft YaHei" if language == "zh" else "Malgun Gothic")
    doc.save(path)


def chunk_paragraphs(paragraphs, max_chars=3000):
    chunks = []
    current = []
    total = 0
    for p in paragraphs:
        extra = len(p) + 1
        if current and total + extra > max_chars:
            chunks.append(current)
            current = []
            total = 0
        current.append(p)
        total += extra
    if current:
        chunks.append(current)
    return chunks


def translate_with_gemini(paragraphs):
    import google.generativeai as genai

    api_key = os.environ.get("GOOGLE_GENERATIVE_AI_API_KEY")
    if not api_key:
        raise RuntimeError("GOOGLE_GENERATIVE_AI_API_KEY is not set")
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(os.environ.get("GEMINI_TRANSLATION_MODEL", "gemini-2.0-flash"))
    translated = []
    chunks = chunk_paragraphs(paragraphs)
    if TRANSLATION_CACHE.exists():
        translated = TRANSLATION_CACHE.read_text(encoding="utf-8").split("\n\n")
        print(f"loaded {len(translated)} cached paragraphs", flush=True)
    for idx, chunk in enumerate(chunks, 1):
        if len(translated) >= sum(len(c) for c in chunks[:idx]):
            continue
        numbered = "\n".join(f"<p id='{i}'>{p}</p>" for i, p in enumerate(chunk))
        prompt = (
            "请把以下韩文研究计划书段落翻译成简体中文。保持原有段落数量和顺序，"
            "保留英文术语、缩写、引用文献、编号、日期和专有名词。不要总结，不要添加解释。"
            "输出格式必须逐段使用 <p id='数字'>译文</p>，id 与输入一致。\n\n"
            f"{numbered}"
        )
        for attempt in range(4):
            try:
                response = model.generate_content(
                    prompt,
                    generation_config={"temperature": 0.1},
                    request_options={"timeout": 120},
                )
                text = response.text
                found = dict(
                    (int(m.group(1)), clean_translation(m.group(2)))
                    for m in re.finditer(r"<p id=['\"]?(\d+)['\"]?>(.*?)</p>", text, re.S)
                )
                if len(found) == len(chunk):
                    translated.extend(found[i] for i in range(len(chunk)))
                    TRANSLATION_CACHE.write_text("\n\n".join(translated), encoding="utf-8")
                    print(f"translated chunk {idx}/{len(chunks)}", flush=True)
                    break
                if attempt == 3:
                    raise RuntimeError(f"chunk {idx}: expected {len(chunk)} paragraphs, got {len(found)}")
            except Exception:
                if attempt == 3:
                    raise
                time.sleep(2 + attempt * 3)
    return translated


def clean_translation(text):
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def main():
    paragraphs = extract_hwp_paragraphs(SOURCE)
    EXTRACTED_TXT.write_text("\n\n".join(paragraphs), encoding="utf-8")
    write_docx(paragraphs, ORIGINAL_DOCX, "ko")
    print(f"wrote {ORIGINAL_DOCX} ({len(paragraphs)} paragraphs)")
    if "--translate" in sys.argv:
        zh = translate_with_gemini(paragraphs)
        write_docx(zh, CHINESE_DOCX, "zh")
        print(f"wrote {CHINESE_DOCX} ({len(zh)} paragraphs)")


if __name__ == "__main__":
    main()
